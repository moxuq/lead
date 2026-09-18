import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..schemas.leads import LeadExport

def export_to_csv(leads: list[LeadExport], filepath: str) -> int:
    dict_leads = [lead.model_dump(by_alias=True) for lead in leads]
    df = pd.DataFrame(dict_leads)
    path = Path(filepath)
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False, encoding='utf-8-sig')
    return len(df)

def export_to_excel(leads: list[LeadExport], filepath: str) -> int:
    dict_leads = [lead.model_dump(by_alias=True) for lead in leads]
    df = pd.DataFrame(dict_leads)
    path = Path(filepath)
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_excel(path, index=False, sheet_name='Leads')
    format_excel(str(path))
    return len(df)

def format_excel(filepath: str):
    wb = load_workbook(filepath)
    ws = wb.active
    ws.freeze_panes = 'A2'
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    wb.save(filepath)
